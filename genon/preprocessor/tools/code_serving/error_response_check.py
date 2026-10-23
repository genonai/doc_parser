#!/usr/bin/env python
"""기본 전처리기 서비스(src/main.py)의 `/run` 으로 오류 응답의 사유 코드·성격과 청크 검증 보고를 확인한다(#528).

파싱 액티비티 처리기와 청킹 액티비티 처리기를 각각 `preprocessor.py` 로 올린 서비스 두 벌을 만든다.
모니모 실물 샘플에서 오류 입력을 만들어 넣고, 응답의 `reason_code`, `error_category`, `error_kind` 가
코드표(`manual/custom_fields_exceptions.md` 의 "사유 코드" 절)와 맞는지 단정한다. 사유 코드는 처리기가
던지는 예외에 실리므로 워커 경로와 같은 분류다. FastAPI TestClient 로 앱을 직접 호출하므로 네트워크가 필요 없다.
#577 전에는 코드서빙 루트 main.py 의 /parser·/chunker 로 확인했다.

설정은 모니모 완성본의 임시 사본을 쓴다. front matter 오류는 현장 설정(on_invalid 기본
ignore)에서 나지 않으므로 사본의 product_slf 에만 `on_invalid: error` 를 켠다.

LLM 오류는 실패를 응답까지 올리는 docling 목차 enrichment(`toc=1`)로 확인한다. custom_fields LLM
추출·문서 요약·표 설명은 strict 여도 경고 후 진행하므로 쓰지 않는다. 텍스트 LLM 프리셋(default)은
정해진 상태로 응답하는 로컬 모의 서버로 돌린다. 목차 경로는 긴 문서를 나눠 부르므로 사전 길이
검사(PARSE_P001)는 여기서 재현하지 않고 단위 테스트가 확인한다.

    genon/preprocessor/.venv/bin/python genon/preprocessor/tools/code_serving/error_response_check.py
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import yaml

REPO_ROOT = Path(__file__).resolve().parents[4]
MONIMO = REPO_ROOT / "genon" / "preprocessor" / "sample_files" / "monimo"
WORK = Path(tempfile.mkdtemp(prefix="error_response_check_"))
_CATEGORY = {"D": "data", "P": "preprocess", "S": "system"}


def _failure_ok(body: dict, want_code: str, want_kind) -> bool:
    """실패 응답이 기대한 사유 코드·성격·유형을 싣는가. 유형은 코드의 분류문자에서 정해진다."""
    return (body.get("code") == 1 and body.get("reason_code") == want_code
            and body.get("error_kind") == want_kind
            and body.get("error_category") == _CATEGORY[want_code.partition("_")[2][:1]])
MAX_CHARS = 1000


def _resource_copy() -> Path:
    """모니모 사이트 완성본 사본(resource/ + activities/) → 사본의 resource/ 경로(GENOS_RESOURCE_DIR).
    처리기는 그 옆 activities/parse·chunk/ 의 설정을 읽는다 — 확인용 변경은 그쪽에 한다.
    product_slf 의 front matter 형식 오류만 실패로 바꾼다."""
    site = WORK / "monimo"
    shutil.copytree(REPO_ROOT / "genon" / "sites" / "monimo", site, ignore=shutil.ignore_patterns("__pycache__"))
    parse, chunk = site / "activities" / "parse", site / "activities" / "chunk"

    def edit(path: Path, change) -> None:
        cfg = yaml.safe_load(path.read_text(encoding="utf-8"))
        change(cfg)
        path.write_text(yaml.safe_dump(cfg, allow_unicode=True, sort_keys=False), encoding="utf-8")

    edit(parse / "custom_field_product_slf.yaml",
         lambda cfg: cfg["source"]["pre"]["markdown"]["front_matter"].__setitem__("on_invalid", "error"))
    # 용어사전은 실물 시트 이름으로 처리 대상 시트를 지정한다(source.sheets).
    edit(parse / "custom_field_term.yaml", lambda cfg: cfg["source"].__setitem__("sheets", ["공통_용어사전"]))
    # 청크 글자 수 상한 보고(max_chars). 운영 값이 정해지지 않아 확인용 값을 넣는다.
    edit(chunk / "config.yaml", lambda cfg: cfg["chunking"]["validation"].__setitem__("max_chars", MAX_CHARS))

    def toc_on(cfg):
        for item in cfg["enrichment"]:
            if "toc" in item:
                # 켜 두되 요청마다 toc 로 끈다(post 의 기본값 toc=0).
                item["toc"]["enable"] = True
    edit(parse / "config.yaml", toc_on)
    return site / "resource"


# 모의 LLM 서버. _LLM_MODE 값으로 응답을 고른다: 상태 번호, "close"(연결 끊기), "slow"(3초 뒤 503).
_LLM_MODE = {"value": 503}


class _MockLLM(BaseHTTPRequestHandler):
    def do_POST(self):  # noqa: N802
        self.rfile.read(int(self.headers.get("Content-Length") or 0))
        mode = _LLM_MODE["value"]
        if mode == "slow":
            time.sleep(3)
            mode = 503
        if mode == "close":
            self.close_connection = True
            self.connection.close()
            return
        status, body = mode, b'{"error": "mock"}'
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Retry-After", "0")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


_SERVER = ThreadingHTTPServer(("127.0.0.1", 0), _MockLLM)
threading.Thread(target=_SERVER.serve_forever, daemon=True).start()
_PRESETS = WORK / "model_presets.local.yaml"
_PRESETS.write_text(yaml.safe_dump({"model_presets": {"default": {
    "url": f"http://127.0.0.1:{_SERVER.server_port}/v1/chat/completions", "api_key": "", "model": "mock"}}}),
    encoding="utf-8")
os.environ["GENOS_MODEL_PRESETS_FILE"] = str(_PRESETS)
os.environ["GENOS_RESOURCE_DIR"] = str(_resource_copy())
# src/main.py 는 common.settings 를 거치는데 그것이 MQ·MinIO 접속 정보를 필수로 요구한다. 이 확인에서는
# 쓰이지 않으므로 더미 값을 채운다. PREPROCESSOR_ID 는 주지 않는다(주면 MinIO 에서 자원을 내려받으려 한다).
for _k, _v in {
    "PROFILE": "dev", "POD_ID": "local",
    "LOG_PATH": json.dumps([tempfile.mkdtemp(prefix="error_response_check_log_")]),
    "MQ_HOST": "127.0.0.1", "MQ_PORT": "5672", "MQ_USER": "guest",
    "MQ_PASSWORD": "guest", "MQ_VHOST": "/", "MQ_EXCHANGE_TYPE": "topic",
    "MQ_EXCHANGE_NAME_LOG": "log", "MQ_QUEUE_NAME_LOG": "log",
    "MQ_QUEUE_BIND_ROUTING_KEY_LOG": "log.#", "MQ_ROUTING_KEY_LOG": "log.local",
    "MINIO_ENDPOINT": "127.0.0.1:9000", "MINIO_ACCESS_KEY": "x", "MINIO_SECRET_KEY": "x",
}.items():
    os.environ.setdefault(_k, _v)
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "genon" / "preprocessor" / "src"))   # src/main.py 의 `from logger import …`


def _service(processor):
    """처리기 하나를 preprocessor.py 로 올린 기본 전처리기 서비스(src/main.py) → TestClient."""
    import importlib.util
    import types

    from fastapi.testclient import TestClient

    module = types.ModuleType("preprocessor")
    module.DocumentProcessor = lambda: processor
    sys.modules["preprocessor"] = module
    spec = importlib.util.spec_from_file_location(
        f"src_main_{id(processor)}", REPO_ROOT / "genon" / "preprocessor" / "src" / "main.py")
    src_main = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(src_main)
    return TestClient(src_main.app)


def _derived_inputs() -> dict:
    """실물 샘플을 조금씩 손상시킨 오류 입력."""
    import openpyxl

    out: dict = {}
    out["empty"] = WORK / "empty.pdf"
    out["empty"].write_bytes(b"")

    faq = (MONIMO / "monimo_faq_json_sample.json").read_text(encoding="utf-8")
    out["broken_json"] = WORK / "monimo_faq_json_sample.json"
    out["broken_json"].write_text(faq[: len(faq) // 2], encoding="utf-8")

    hpp = json.loads((MONIMO / "monimo_product_hpp_sample.json").read_text(encoding="utf-8"))
    hpp["cardTitle"] = ""
    out["hpp_no_title"] = WORK / "monimo_product_hpp_sample.json"
    out["hpp_no_title"].write_text(json.dumps(hpp, ensure_ascii=False), encoding="utf-8")

    # 6행이 열 이름 행이다. '서비스 개요' 를 '메뉴중분류' 로 바꿔 정규화 후 중복을 만든다.
    wb = openpyxl.load_workbook(MONIMO / "monimo_menu_sample.xlsx")
    ws = wb.active
    for cell in ws[6]:
        if cell.value == "서비스 개요":
            cell.value = "메뉴중분류"
    out["menu_dup_header"] = WORK / "monimo_menu_sample.xlsx"
    wb.save(out["menu_dup_header"])

    slf = (MONIMO / "monimo_product_slf_sample.md").read_text(encoding="utf-8")
    out["slf_dup_key"] = WORK / "dup_key" / "monimo_product_slf_sample.md"
    out["slf_dup_key"].parent.mkdir()
    out["slf_dup_key"].write_text(slf.replace("---\n", "---\nauthor: \"중복\"\n", 1), encoding="utf-8")

    # `## 원문 전사` 뒤에 제외 대상 섹션만 남겨 섹션 선택 후 본문이 없게 만든다.
    sections = (MONIMO / "monimo_product_slf_sections_sample.md").read_text(encoding="utf-8")
    head = sections.split("## 원문 1쪽", 1)[0]
    out["slf_no_body"] = WORK / "no_body" / "monimo_product_slf_sections_sample.md"
    out["slf_no_body"].parent.mkdir()
    out["slf_no_body"].write_text(head, encoding="utf-8")

    # 용어사전 시트 뒤에 메뉴 시트를 붙인 통합 문서. 메뉴 시트에는 용어사전 필수 열이 없다.
    book = openpyxl.load_workbook(MONIMO / "monimo_term_sample.xlsx")
    menu = openpyxl.load_workbook(MONIMO / "monimo_menu_sample.xlsx").active
    extra = book.create_sheet(menu.title)
    for row in menu.iter_rows(values_only=True):
        extra.append(row)
    out["term_plus_menu"] = WORK / "monimo_term_plus_menu.xlsx"
    book.save(out["term_plus_menu"])

    # 제외할 시트에 완전히 같은 열 이름이 둘 있는 통합 문서. 지정한 시트 처리를 막지 않아야 한다.
    book = openpyxl.load_workbook(MONIMO / "monimo_term_sample.xlsx")
    broken = book.create_sheet("메모")
    broken.append(["비고", "비고"])
    broken.append(["a", "b"])
    out["term_plus_dup"] = WORK / "monimo_term_plus_dup.xlsx"
    book.save(out["term_plus_dup"])

    # 지정한 용어사전 시트에 열 이름 행만 남긴 문서(이번 배치 데이터 0건).
    book = openpyxl.load_workbook(MONIMO / "monimo_term_sample.xlsx")
    sheet = book["공통_용어사전"]
    header_row = next(i for i, row in enumerate(sheet.iter_rows(values_only=True), 1)
                      if "용어(국문)" in row)
    sheet.delete_rows(header_row + 1, sheet.max_row)
    out["term_header_only"] = WORK / "monimo_term_header_only.xlsx"
    book.save(out["term_header_only"])

    # 읽을 수 없는 front matter 원천. 원천 형식 오류(D005)가 아니라 파일 오류(S003)여야 한다.
    out["locked_md"] = WORK / "locked" / "monimo_product_slf_sample.md"
    out["locked_md"].parent.mkdir()
    out["locked_md"].write_text(slf, encoding="utf-8")
    out["locked_md"].chmod(0)

    out["locked_json"] = WORK / "locked.json"
    out["locked_json"].write_text("{}", encoding="utf-8")
    out["locked_json"].chmod(0)
    return out


def main() -> int:
    from genon.preprocessor.activities import _folders
    from genon.preprocessor.activities.chunk import DocumentProcessor as ChunkProcessor
    from genon.preprocessor.activities.parse import DocumentProcessor as ParseProcessor

    clients = {"/parser": _service(ParseProcessor(config_path=str(_folders.config_path("parse")))),
               "/chunker": _service(ChunkProcessor(config_path=str(_folders.config_path("chunk"))))}
    src = _derived_inputs()

    def post(route: str, file_path, params: dict) -> dict:
        """route 는 처리기 이름표다 — /parser 는 파싱 처리기를, /chunker 는 청킹 처리기를 올린 서비스의 /run."""
        if route == "/parser":
            params = {"toc": 0, **params}
        return clients[route].post("/run", json={"file_path": str(file_path), "params": params}).json()

    parsed = post("/parser", MONIMO / "monimo_product_ssf_sample.md", {"doc_type": "product_ssf"})
    assert parsed.get("code") == 0, parsed.get("errMsg")
    doc = parsed["data"]
    corrupt = {**doc, "document": {**doc["document"], "texts": "broken"}}

    # (설명, 경로, file_path, params, 기대 reason_code, 기대 error_kind)
    cases = [
        ("필수 열 부재(term 엑셀을 menu 로)", "/parser", MONIMO / "monimo_term_sample.xlsx",
         {"doc_type": "menu"}, "PARSE_D001", "permanent"),
        ("records 키 부재(cs_slf JSON 을 faq 로)", "/parser", MONIMO / "monimo_cs_slf_sample.json",
         {"doc_type": "faq"}, "PARSE_D002", "permanent"),
        ("필수 공통 필드 공백(상품명 비움)", "/parser", src["hpp_no_title"],
         {"doc_type": "product_hpp"}, "PARSE_D003", "permanent"),
        ("열 이름 중복(정규화 후)", "/parser", src["menu_dup_header"],
         {"doc_type": "menu"}, "PARSE_D004", "permanent"),
        ("front matter 중복 키", "/parser", src["slf_dup_key"],
         {"doc_type": "product_slf"}, "PARSE_D005", "permanent"),
        ("front matter 원천 권한 거부", "/parser", src["locked_md"],
         {"doc_type": "product_slf"}, "PARSE_S003", "permanent"),
        ("빈 파일", "/parser", src["empty"], {"doc_type": "product_ssf"}, "PARSE_D009", "permanent"),
        ("JSON 문법 손상(faq 절반)", "/parser", src["broken_json"],
         {"doc_type": "faq"}, "PARSE_D010", "permanent"),
        ("구분자 레코드 없음(md 를 cs_ssf 로)", "/parser", MONIMO / "monimo_product_ssf_sample.md",
         {"doc_type": "cs_ssf"}, "PARSE_D011", "permanent"),
        ("섹션 선택 후 본문 없음", "/parser", src["slf_no_body"],
         {"doc_type": "product_slf"}, "PARSE_D012", "permanent"),
        ("지정한 시트 없음(메뉴 엑셀을 term 으로)", "/parser", MONIMO / "monimo_menu_sample.xlsx",
         {"doc_type": "term"}, "PARSE_D013", "permanent"),
        ("미등록 doc_type", "/parser", MONIMO / "monimo_product_ssf_sample.md",
         {"doc_type": "no_such_type"}, "PARSE_P002", "permanent"),
        ("청커 입력 누락", "/chunker", "", {}, "CHUNK_P002", "permanent"),
        ("파서 결과가 아닌 입력", "/chunker", "", {"document": {"foo": 1}}, "CHUNK_P003", "permanent"),
        ("파서 결과 손상(실물 산출 texts 손상)", "/chunker", "", {"document": corrupt},
         "CHUNK_P004", "permanent"),
        ("파서 결과 JSON 문법 손상", "/chunker", src["broken_json"], {}, "CHUNK_P004", "permanent"),
        ("입력 파일 권한 거부", "/chunker", src["locked_json"], {}, "CHUNK_S003", "permanent"),
    ]

    # chmod 로 만든 권한 거부는 root 로 실행하면 재현되지 않는다. 그때는 해당 사례를 뺀다.
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        print("[SKIP] root 로 실행 중이라 권한 거부 사례를 뺍니다.")
        cases = [case for case in cases if case[4] not in ("PARSE_S003", "CHUNK_S003")]

    failures = 0
    for label, route, file_path, params, want_code, want_kind in cases:
        body = post(route, file_path, params)
        got = (body.get("reason_code"), body.get("error_kind"))
        ok = _failure_ok(body, want_code, want_kind)
        failures += not ok
        print(f"[{'PASS' if ok else 'FAIL'}] {label}: reason_code={got[0]} "
              f"error_category={body.get('error_category')} error_kind={got[1]}"
              + ("" if ok else f"  기대={want_code}/{want_kind}  errMsg={str(body.get('errMsg'))[:120]}"))

    small = MONIMO / "monimo_product_ssf_sample.md"
    llm_cases = [
        (503, small, "product_ssf", "PARSE_S101", "transient"),
        (429, small, "product_ssf", "PARSE_S101", "transient"),
        ("close", small, "product_ssf", "PARSE_S101", "transient"),
        (401, small, "product_ssf", "PARSE_S001", "permanent"),
    ]
    for mode, path, doc_type, want_code, want_kind in llm_cases:
        _LLM_MODE["value"] = mode
        body = post("/parser", path, {"doc_type": doc_type, "toc": 1})
        got = (body.get("reason_code"), body.get("error_kind"))
        ok = _failure_ok(body, want_code, want_kind) and body.get("stage") == "enrichment"
        failures += not ok
        print(f"[{'PASS' if ok else 'FAIL'}] 목차 LLM {mode}({path.name}): reason_code={got[0]} error_kind={got[1]} "
              f"stage={body.get('stage')}"
              + ("" if ok else f"  기대={want_code}/{want_kind}  errMsg={str(body.get('errMsg'))[:120]}"))
    cases += llm_cases

    # 시트 지정: 메뉴 시트가 붙은 통합 문서도 지정한 용어사전 시트만 처리해 성공한다.
    alone = post("/parser", MONIMO / "monimo_term_sample.xlsx", {"doc_type": "term"})
    mixed = post("/parser", src["term_plus_menu"], {"doc_type": "term"})
    counts = [len((b.get("data") or {}).get("elements") or []) for b in (alone, mixed)]
    ok = alone.get("code") == mixed.get("code") == 0 and counts[0] == counts[1] > 0
    failures += not ok
    print(f"[{'PASS' if ok else 'FAIL'}] 시트 지정(용어사전+메뉴 통합 문서): elements={counts} "
          f"errMsg={str(mixed.get('errMsg'))[:80]}")
    cases.append("sheets")
    for label, path in [("제외한 시트의 열 이름 중복은 무시", src["term_plus_dup"]),
                        ("지정한 시트에 데이터 0건이어도 시트 없음이 아님", src["term_header_only"])]:
        body = post("/parser", path, {"doc_type": "term"})
        ok = body.get("code") == 0
        failures += not ok
        print(f"[{'PASS' if ok else 'FAIL'}] {label}: code={body.get('code')} "
              f"elements={len((body.get('data') or {}).get('elements') or [])} "
              f"errMsg={str(body.get('errMsg'))[:100]}")
        cases.append(label)

    # 청킹 중 빠진 메타 필드 보고(missing_fields). 문서는 실패하지 않는다.
    # 실물 파서 산출에 외부 payload 처럼 예약 이름의 메타를 싣는다.
    faq = post("/parser", MONIMO / "monimo_faq_json_sample.json", {"doc_type": "faq"})["data"]
    for element in faq.get("elements") or []:
        element.setdefault("metadata", {})["i_page"] = 99
    slf = post("/parser", MONIMO / "monimo_product_slf_sample.md", {"doc_type": "product_slf"})["data"]
    # 문서 메타는 docling 문서의 key_value_items 로 청커에 넘어간다. 실물 항목에 키·값 셀을 덧붙인다.
    cells = slf["document"]["key_value_items"][0]["graph"]["cells"]
    for label, text in (("key", "file_path"), ("value", "원본.pdf")):
        cells.append({"label": label, "cell_id": len(cells), "text": text, "orig": text,
                      "prov": None, "item_ref": None})
    for label, doc_type, document, want in [
        ("행 메타의 예약 키가 덮어써짐(faq)", "faq", faq, "i_page"),
        ("문서 메타의 예약 키가 빠짐(product_slf)", "product_slf", slf, "file_path"),
        ("정상 산출은 보고 없음(product_slf)", "product_slf",
         post("/parser", MONIMO / "monimo_product_slf_sample.md", {"doc_type": "product_slf"})["data"], None),
    ]:
        body = post("/chunker", "", {"doc_type": doc_type, "document": document})
        missing = (body.get("chunk_validation") or {}).get("missing_fields") or {}
        ok = body.get("code") == 0 and (want in missing if want else not missing)
        failures += not ok
        print(f"[{'PASS' if ok else 'FAIL'}] {label}: missing_fields={missing} code={body.get('code')}")
        cases.append(label)

    # 글자 수 상한 보고(max_length). 보고 건수가 실제로 상한을 넘은 출력 청크 수와 같아야 한다.
    hpp = MONIMO / "monimo_product_hpp_rich_table_sample.json"
    parsed_hpp = post("/parser", hpp, {"doc_type": "product_hpp"})["data"]
    body = post("/chunker", "", {"doc_type": "product_hpp", "document": parsed_hpp})
    report = (body.get("chunk_validation") or {}).get("max_length") or {}
    over = sum(len(v.get("text") or "") > MAX_CHARS for v in body.get("data") or [])
    ok = (body.get("code") == 0 and over > 0 and report.get("count") == over
          and report.get("limit") == MAX_CHARS
          and all(item["n_char"] > MAX_CHARS for item in report.get("items") or []))
    failures += not ok
    print(f"[{'PASS' if ok else 'FAIL'}] 글자 수 상한 {MAX_CHARS} 초과 보고(product_hpp 표): "
          f"보고={report.get('count')} 실제={over} 청크={len(body.get('data') or [])}")
    cases.append("max_length")

    # 요청 deadline 초과. 처리기 안에서 난 시간 초과와 구분되어 COMMON_S102 로 실린다.
    # 동기 변환 구간은 deadline 이 끊지 못하므로(알려진 한계) 비동기 LLM 호출을 기다리는 동안 넘긴다.
    _LLM_MODE["value"] = "slow"
    body = post("/parser", REPO_ROOT / "genon" / "preprocessor" / "sample_files" / "json" /
                "monimo_card_sample.json", {"doc_type": "card", "request_deadline": 1})
    _LLM_MODE["value"] = 503
    ok = _failure_ok(body, "COMMON_S102", "timeout") and body.get("stage") == "request"
    failures += not ok
    print(f"[{'PASS' if ok else 'FAIL'}] 요청 deadline 초과: reason_code={body.get('reason_code')} "
          f"error_kind={body.get('error_kind')} stage={body.get('stage')} code={body.get('code')}")
    cases.append("deadline")

    # 요청 형식 오류는 라우트 함수에 닿기 전에 FastAPI 가 거부한다.
    body = clients["/parser"].post("/run", json={"params": {}}).json()
    ok = body.get("reason_code") == "COMMON_P002"
    failures += not ok
    print(f"[{'PASS' if ok else 'FAIL'}] 요청 형식 오류(file_path 누락): reason_code={body.get('reason_code')}")

    print(f"\n{len(cases) + 1 - failures}/{len(cases) + 1} 통과")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    finally:
        for locked in (WORK / "locked.json", WORK / "locked" / "monimo_product_slf_sample.md"):
            if locked.exists():
                locked.chmod(0o600)
        shutil.rmtree(WORK, ignore_errors=True)
